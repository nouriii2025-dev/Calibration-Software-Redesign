from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from django.conf import settings
from django.http import FileResponse
from decimal import Decimal


EXCEL_HEADERS = [
    "Scope of Work",
    "Datasheet",
    "Job No",
    "Cert. No",
    "Customer name",
    "Address",
    "Tel",
    "FAX",
    "PO",
    "Instrument",
    "Range",
    "Model",
    "Make",
    "Serial",
    "Cal by",
    "Cal. Date",
    "Due Date",
    "Resolution",
    "Accuracy",
    "Tag",
    "Location",
    "Temp",
    "Rel. Humidity",
    "Atm. Press",
    "Ratio",
    "Issue Date",
    "Remarks Status",
]


EXCEL_DIR = Path(settings.MEDIA_ROOT) / "exports"
EXCEL_FILE = EXCEL_DIR / "calibration_register.xlsx"


def get_excel_file():
    """
    Create the master Excel file if it does not already exist.
    """

    EXCEL_DIR.mkdir(parents=True, exist_ok=True)

    if not EXCEL_FILE.exists():

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Calibration Register"

        for column, header in enumerate(EXCEL_HEADERS, start=1):

            cell = worksheet.cell(
                row=1,
                column=column,
                value=header
            )

            cell.font = Font(
                bold=True,
                name="Arial"
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="9DC3E6"
            )

        worksheet.freeze_panes = "A2"

        # Give the columns reasonable widths.
        widths = [
            20, 18, 14, 14, 25, 30, 15, 15,
            15, 25, 20, 20, 20, 20, 20, 15,
            15, 15, 15, 18, 18, 15, 18, 18,
            15, 15, 20
        ]

        for index, width in enumerate(widths, start=1):
            worksheet.column_dimensions[
                worksheet.cell(row=1, column=index).column_letter
            ].width = width

        workbook.save(EXCEL_FILE)

    return EXCEL_FILE


def get_certificate_status(certificate):
    """
    Determine the overall certificate status from its result rows.
    """

    results = certificate.results.all()

    if not results.exists():
        return ""

    remarks = []

    for result in results:

        value = (result.remarks or "").strip().upper()

        if value:
            remarks.append(value)

    if not remarks:
        return ""

    if "FAIL" in remarks:
        return "FAIL"

    if all(value == "PASS" for value in remarks):
        return "PASS"

    return ", ".join(sorted(set(remarks)))


def certificate_to_excel_row(certificate):
    """
    Convert one Django Certificate into one Excel row.
    """

    job = certificate.job
    line_item = certificate.line_item

    instrument = ""
    assigned_technician = ""

    if line_item:

        if line_item.instrument:
            instrument = str(line_item.instrument)

        if line_item.assigned_to:
            assigned_technician = str(
                line_item.assigned_to
            )

    # ---------------------------------------------------------
    # Range
    # ---------------------------------------------------------

    range_value = ""

    if line_item:

        if (
            line_item.range_from is not None
            and line_item.range_to is not None
        ):
            range_value = (
                f"{line_item.range_from:g} "
                f"to "
                f"{line_item.range_to:g}"
            )

        elif line_item.range_to is not None:
            range_value = str(line_item.range_to)

        elif line_item.range_from is not None:
            range_value = str(line_item.range_from)

    # ---------------------------------------------------------
    # Accuracy
    # ---------------------------------------------------------

    accuracy = ""

    if certificate.device_accuracy:

        accuracy = str(
            certificate.device_accuracy
        )

        if certificate.device_accuracy_unit:
            accuracy += f" {certificate.device_accuracy_unit}"

    status = get_certificate_status(certificate)

    return [
        job.scope_of_work,                      # Scope of Work
        certificate.datasheet_number,           # Datasheet
        job.job_number,                         # Job No
        certificate.number,                     # Cert. No
        job.customer_name,                      # Customer name
        job.customer_address,                   # Address
        job.telephone_number,                   # Tel
        job.fax_number,                         # FAX
        job.po_number,                          # PO
        instrument,                             # Instrument
        range_value,                            # Range
        line_item.model if line_item else "",   # Model
        "",                                     # Make
        certificate.device_serial,              # Serial
        certificate.calibrated_by,              # Cal by
        certificate.calibration_date,           # Cal. Date
        certificate.re_calibration_date,        # Due Date
        certificate.device_resolution,          # Resolution
        accuracy,                               # Accuracy
        certificate.device_tag_number,          # Tag
        job.location,                           # Location
        certificate.lab_temperature,            # Temp
        certificate.lab_humidity,               # Rel. Humidity
        certificate.ambient_pressure,           # Atm. Press
        certificate.device_ratio,               # Ratio
        certificate.issue_date,                 # Issue Date
        status,                                 # Remarks Status
    ]


# def export_certificate_to_excel(certificate):
#     """
#     Add/update one certificate in the master Excel workbook.

#     The certificate number is used as the unique key.
#     Therefore exporting the same certificate multiple times
#     does not create duplicate rows.
#     """

#     excel_file = get_excel_file()

#     workbook = load_workbook(excel_file)

#     if "Calibration Register" in workbook.sheetnames:
#         worksheet = workbook["Calibration Register"]
#     else:
#         worksheet = workbook.create_sheet(
#             "Calibration Register"
#         )

#     # ---------------------------------------------------------
#     # Make sure headers exist
#     # ---------------------------------------------------------

#     if worksheet.max_row == 0:

#         for column, header in enumerate(
#             EXCEL_HEADERS,
#             start=1
#         ):
#             worksheet.cell(
#                 row=1,
#                 column=column,
#                 value=header
#             )

#     # ---------------------------------------------------------
#     # Find existing certificate row
#     # ---------------------------------------------------------

#     certificate_column = EXCEL_HEADERS.index(
#         "Cert. No"
#     ) + 1

#     existing_row = None

#     for row in range(2, worksheet.max_row + 1):

#         cell_value = worksheet.cell(
#             row=row,
#             column=certificate_column
#         ).value

#         if str(cell_value).strip() == str(
#             certificate.number
#         ).strip():

#             existing_row = row
#             break

#     # ---------------------------------------------------------
#     # If certificate already exists:
#     # UPDATE it.
#     #
#     # Otherwise:
#     # ADD a new row.
#     # ---------------------------------------------------------

#     if existing_row is None:

#         row_number = worksheet.max_row + 1

#     else:

#         row_number = existing_row

#     values = certificate_to_excel_row(
#         certificate
#     )

#     for column, value in enumerate(
#         values,
#         start=1
#     ):

#         cell = worksheet.cell(
#             row=row_number,
#             column=column,
#             value=value
#         )

#         cell.alignment = Alignment(
#             vertical="center"
#         )

#     # ---------------------------------------------------------
#     # Date formatting
#     # ---------------------------------------------------------

#     date_headers = {
#         "Cal. Date",
#         "Due Date",
#         "Issue Date",
#     }

#     for column, header in enumerate(
#         EXCEL_HEADERS,
#         start=1
#     ):

#         if header in date_headers:

#             worksheet.cell(
#                 row=row_number,
#                 column=column
#             ).number_format = "dd/mm/yyyy"

#     # ---------------------------------------------------------
#     # Add borders to data row
#     # ---------------------------------------------------------

#     thin_border = Border(
#         left=Side(style="thin"),
#         right=Side(style="thin"),
#         top=Side(style="thin"),
#         bottom=Side(style="thin"),
#     )

#     for column in range(
#         1,
#         len(EXCEL_HEADERS) + 1
#     ):

#         worksheet.cell(
#             row=row_number,
#             column=column
#         ).border = thin_border

#     # ---------------------------------------------------------
#     # Save master workbook
#     # ---------------------------------------------------------

#     workbook.save(excel_file)

#     return excel_file

def export_certificate_to_excel(certificate):
    """
    Add or update a certificate in the single master Excel workbook.

    All certificates are stored in:
        MEDIA_ROOT/exports/calibration_register.xlsx

    If the certificate already exists, its row is updated.
    If it does not exist, a new row is appended.
    """

    excel_file = get_excel_file()

    # ---------------------------------------------------------
    # Open the existing master workbook
    # ---------------------------------------------------------

    workbook = load_workbook(excel_file)

    # ---------------------------------------------------------
    # Get the Calibration Register sheet
    # ---------------------------------------------------------

    if "Calibration Register" in workbook.sheetnames:
        worksheet = workbook["Calibration Register"]
    else:
        worksheet = workbook.create_sheet("Calibration Register")

        for column, header in enumerate(EXCEL_HEADERS, start=1):
            cell = worksheet.cell(
                row=1,
                column=column,
                value=header
            )

            cell.font = Font(
                bold=True,
                name="Arial"
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="9DC3E6"
            )

        worksheet.freeze_panes = "A2"

    # ---------------------------------------------------------
    # Make sure headers exist
    # ---------------------------------------------------------

    for column, header in enumerate(EXCEL_HEADERS, start=1):

        if worksheet.cell(row=1, column=column).value != header:
            worksheet.cell(
                row=1,
                column=column,
                value=header
            )

            worksheet.cell(
                row=1,
                column=column
            ).font = Font(
                bold=True,
                name="Arial"
            )

    # ---------------------------------------------------------
    # Certificate number is our unique key
    # ---------------------------------------------------------

    certificate_column = EXCEL_HEADERS.index("Cert. No") + 1

    existing_row = None

    for row in range(2, worksheet.max_row + 1):

        existing_certificate = worksheet.cell(
            row=row,
            column=certificate_column
        ).value

        if (
            existing_certificate is not None
            and str(existing_certificate).strip()
            == str(certificate.number).strip()
        ):
            existing_row = row
            break

    # ---------------------------------------------------------
    # Existing certificate -> update
    # New certificate -> append
    # ---------------------------------------------------------

    if existing_row is not None:
        row_number = existing_row
    else:
        row_number = worksheet.max_row + 1

    # ---------------------------------------------------------
    # Convert certificate to Excel data
    # ---------------------------------------------------------

    values = certificate_to_excel_row(certificate)

    # ---------------------------------------------------------
    # Write the certificate data
    # ---------------------------------------------------------

    for column, value in enumerate(values, start=1):

        cell = worksheet.cell(
            row=row_number,
            column=column,
            value=value
        )

        cell.alignment = Alignment(
            vertical="center"
        )

    # ---------------------------------------------------------
    # Date formatting
    # ---------------------------------------------------------

    date_headers = {
        "Cal. Date",
        "Due Date",
        "Issue Date",
    }

    for column, header in enumerate(EXCEL_HEADERS, start=1):

        if header in date_headers:

            worksheet.cell(
                row=row_number,
                column=column
            ).number_format = "dd/mm/yyyy"

    # ---------------------------------------------------------
    # Borders
    # ---------------------------------------------------------

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for column in range(1, len(EXCEL_HEADERS) + 1):

        worksheet.cell(
            row=row_number,
            column=column
        ).border = thin_border

    # ---------------------------------------------------------
    # Save the SAME workbook
    # ---------------------------------------------------------

    workbook.save(excel_file)

    return excel_file